#!/usr/bin/env python3
"""시즌 진행현황 시트(xlsx 내보내기) → 사이트용 시즌별 암호화 데이터.

사용법
  python3 tools/build_data.py --xlsx 진행현황.xlsx --out data --password-file ~/.gsb/site.json
  python3 tools/build_data.py --xlsx 진행현황.xlsx --out data --password '...'   (비추천: 기록에 남음)
  옵션 --plain  : 암호화하지 않은 사본(data_plain/)도 같이 만든다. 확인용, 저장소에 올리지 않는다.

동작
  - 시트의 `보드_{시즌} 신규 / 리오더 / 드랍` 탭을 읽어 시즌별 파일 data/seasons/{시즌}.json 을 만든다.
  - 시트에 없는 시즌(예전 시즌)은 기존 파일을 그대로 둔다. 그래서 시즌이 쌓인다.
  - 모든 파일은 AES-GCM(키: PBKDF2-SHA256, 비밀번호)으로 암호화한다. 사이트는 같은 비밀번호로 연다.
  - 원가·판가는 이 시트에 없고, 넣지도 않는다.
"""
import argparse, base64, datetime as dt, json, os, re, sys
from pathlib import Path

import openpyxl
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
from cryptography.hazmat.primitives import hashes

ITER = 250_000
CHECK_TEXT = "renacts-ok"
TAB_RE = re.compile(r"^보드_(\d{2}(?:SS|FW))\s*(신규|리오더|드랍)$")
FIELDS = {  # 시트 머리글 → 데이터 키
    "시즌": "season", "구분": "type", "품번": "code", "컬러": "color",
    "아이템명 (영문)": "name", "기획명": "plan", "현재 단계": "stage", "상태": "status",
    "업데이트 날짜": "updated", "최근 내용": "recent", "메모": "memo",
}
DEFAULT_STAGES = ["01_개발진행중", "02_패턴", "03_프로토샘플", "04_QC · 재QC", "05_APP", "06_메인수납", "07_종결", "08_드랍"]
KST = dt.timezone(dt.timedelta(hours=9))


def b64(b): return base64.b64encode(b).decode()
def unb64(s): return base64.b64decode(s)


def derive(password, salt):
    return PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=salt, iterations=ITER).derive(password.encode())


def enc(key, obj):
    iv = os.urandom(12)
    data = json.dumps(obj, ensure_ascii=False, separators=(",", ":")).encode()
    return {"v": 1, "iv": b64(iv), "ct": b64(AESGCM(key).encrypt(iv, data, None))}


def dec(key, box):
    return json.loads(AESGCM(key).decrypt(unb64(box["iv"]), unb64(box["ct"]), None))


def cell(v):
    if v is None: return ""
    if isinstance(v, (dt.datetime, dt.date)): return v.strftime("%Y-%m-%d")
    return str(v).strip()


def season_key(label):
    m = re.match(r"(\d{2})(SS|FW)", label)
    return (int(m.group(1)), 1 if m.group(2) == "FW" else 0) if m else (0, 0)


def read_sheet(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    stages = DEFAULT_STAGES
    if "설정" in wb.sheetnames:
        col = [cell(r[0]) for r in wb["설정"].iter_rows(min_row=2, values_only=True)]
        col = [c for c in col if c]
        if col: stages = col
    seasons = {}
    for ws in wb:
        m = TAB_RE.match(ws.title.strip())
        if not m: continue
        label, tab_type = m.groups()
        rows = list(ws.iter_rows(values_only=True))
        if not rows: continue
        hdr = [cell(h) for h in rows[0]]
        for r in rows[1:]:
            o = {}
            for i, h in enumerate(hdr):
                if h in FIELDS: o[FIELDS[h]] = cell(r[i]) if i < len(r) else ""
            if not (o.get("code") or o.get("plan") or o.get("name")):
                continue  # 빈 줄 · 양식 줄
            o["season"] = o.get("season") or label
            o["type"] = o.get("type") or tab_type
            o["tab"] = tab_type
            if tab_type == "드랍" and not o.get("stage"): o["stage"] = "08_드랍"
            seasons.setdefault(label, []).append(o)
        seasons.setdefault(label, seasons.get(label, []))
    return stages, seasons


QSTAGES = ["작업 지시", "재단 수량", "워싱처 입고", "완성집 입고", "수납용 입고", "3PL 입고"]
QKEYS = ["order", "cut", "wash", "finish", "store", "tpl"]
INFO_KEYS = {"품번": "code", "영문 판매명 (Flow FIX)": "name", "기획명": "plan", "시즌": "season", "구분": "type", "차수": "round",
             "카테고리": "category", "상태": "status", "현재 단계": "stage", "최근 일자": "updated", "최근 내용": "recent",
             "혼용율 · 겉감": "mixOuter", "혼용율 · 안감": "mixLining", "혼용율 · 배색(아더)": "mixOther", "원산지 (제조국)": "origin",
             "봉제공장": "factory", "FIT": "fit", "원단": "fabric", "규격 · 중량": "spec", "워싱 여부": "wash", "워싱기법": "washMethod",
             "사이즈 구성": "sizes", "세트 연결": "set", "위키 문서": "wiki", "메모": "memo"}
DUMMY_BARCODE = "0000000000000"


def num(v):
    if v in (None, ""): return None
    try:
        f = float(str(v).replace(",", ""))
        return int(f) if f == int(f) else f
    except ValueError:
        return None


def is_product(wb):
    return all(s in wb.sheetnames for s in ("기본정보", "수량", "히스토리"))


def parse_product(path):
    """제품 파일 양식 v3 (파일 하나 = 품번 1건, 컬러 1~5). 원가 탭은 읽지 않는다."""
    with open(path, "rb") as fh:  # 확장자가 없는 파일도 읽도록 파일 객체로 연다
        wb = openpyxl.load_workbook(fh, data_only=True)
    if not is_product(wb): return None
    o, colors = {}, []
    for r in wb["기본정보"].iter_rows(min_row=1, values_only=True):
        lab = cell(r[1]) if len(r) > 1 else ""
        if re.fullmatch(r"컬러 \d", lab):
            code = cell(r[2]); name = cell(r[3]) if len(r) > 3 else ""
            if code: colors.append({"code": code, "name": name})
        elif lab in INFO_KEYS:
            o[INFO_KEYS[lab]] = cell(r[2])
    if not o.get("code"): return None
    # 히스토리
    hist = []
    for r in wb["히스토리"].iter_rows(min_row=5, values_only=True):
        d, col, st, txt, src, link = (list(r) + [None] * 7)[1:7]
        if not (d or st or txt): continue
        hist.append({"date": cell(d), "color": cell(col), "stage": cell(st), "text": cell(txt), "source": cell(src), "link": cell(link)})
    # 수량 (컬러 블록)
    qty = []
    rows = list(wb["수량"].iter_rows(values_only=True))
    for i, r in enumerate(rows):
        lab = cell(r[1]) if len(r) > 1 else ""
        m = re.match(r"컬러 (\d)", lab)
        if not m or "비어 있음" in lab: continue
        k = int(m.group(1)) - 1
        block = {"color": colors[k]["code"] if k < len(colors) else lab, "stages": {}}
        for j, key in enumerate(QKEYS):
            rr = rows[i + 1 + j]
            sizes = [num(v) for v in rr[2:7]]
            if any(v is not None for v in sizes):
                block["stages"][key] = {"sizes": sizes, "total": num(rr[7])}
        dr = rows[i + 7]
        diff = [num(v) for v in dr[2:7]]
        if any(v is not None for v in diff): block["diff"] = {"sizes": diff, "total": num(dr[7])}
        qty.append(block)
    # 바코드
    barcodes = []
    for r in wb["바코드"].iter_rows(min_row=5, values_only=True):
        col, size, bc, kind, reg, memo = (list(r) + [None] * 8)[1:7]
        bc = cell(bc)
        if not bc or bc == DUMMY_BARCODE or "예시" in cell(memo): continue
        barcodes.append({"color": cell(col), "size": cell(size), "barcode": bc, "kind": cell(kind), "date": cell(reg), "memo": cell(memo)})
    # 3PL 검수
    ws = wb["3PL검수"]
    tpl = {"date": cell(ws["C4"].value), "vendor": cell(ws["F4"].value), "rows": []}
    for r in ws.iter_rows(min_row=7, max_row=46, values_only=True):
        col, size, bc, pname, pack, insp, diff, res = (list(r) + [None] * 10)[1:9]
        bc = cell(bc)
        if not bc or bc == DUMMY_BARCODE: continue
        tpl["rows"].append({"color": cell(col), "size": cell(size), "barcode": bc, "packed": num(pack), "checked": num(insp), "diff": num(diff), "result": cell(res)})
    item = {k: o.get(k, "") for k in ("season", "type", "code", "name", "plan", "stage", "status", "updated", "recent", "memo", "category", "round")}
    item["color"] = ", ".join(c["code"] for c in colors)
    item["tab"] = "드랍" if item["status"] == "드랍" else (item["type"] or "신규")
    if item["tab"] == "드랍" and not item["stage"]: item["stage"] = "08_드랍"
    item["source"] = "product"
    item["file"] = Path(path).name
    item["detail"] = {"colors": colors, "info": {k: o.get(k, "") for k in ("mixOuter", "mixLining", "mixOther", "origin", "factory", "fit", "fabric", "spec", "wash", "washMethod", "sizes", "set", "wiki")},
                      "history": hist, "qty": qty, "barcodes": barcodes, "tpl": tpl}
    return item


def read_products(folder):
    found = []
    for p in sorted(Path(folder).rglob("*")):
        if not p.is_file() or p.name.startswith(("~$", ".")): continue
        try:
            it = parse_product(p)
        except Exception as e:  # 엑셀이 아닌 파일 등
            continue
        if it:
            if not it["category"]:
                m = re.match(r"^[\d._\s]*(.+)$", p.parent.name); it["category"] = m.group(1).strip() if m else ""
            found.append(it)
    return found


def load_password(args):
    if args.password: return args.password
    if args.password_file:
        p = Path(os.path.expanduser(args.password_file))
        j = json.loads(p.read_text(encoding="utf-8"))
        return j["sitePassword"]
    pw = os.environ.get("SITE_PASSWORD")
    if pw: return pw
    sys.exit("비밀번호가 없습니다. --password-file 또는 SITE_PASSWORD 를 주세요.")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True, help="시즌 진행현황 시트 (xlsx 내보내기)")
    ap.add_argument("--products", help="제품 파일 폴더 (하위 폴더까지 읽음, 품번 1파일)")
    ap.add_argument("--out", default="data")
    ap.add_argument("--password"); ap.add_argument("--password-file")
    ap.add_argument("--sheet-url", default="")
    ap.add_argument("--new-password", action="store_true", help="비밀번호를 바꾼다(모든 시즌 파일을 새 비밀번호로 다시 암호화)")
    ap.add_argument("--plain", action="store_true")
    args = ap.parse_args()

    password = load_password(args)
    out = Path(args.out); (out / "seasons").mkdir(parents=True, exist_ok=True)
    man_path = out / "manifest.json"
    old = json.loads(man_path.read_text(encoding="utf-8")) if man_path.exists() else None

    # 이전 시즌 보존을 위해 기존 키로 옛 파일을 풀어 둔다
    keep = {}
    if old:
        old_key = derive(password, unb64(old["kdf"]["salt"])) if not args.new_password else None
        if not args.new_password:
            try:
                assert dec(old_key, old["check"]) == CHECK_TEXT
            except Exception:
                sys.exit("비밀번호가 기존 데이터와 다릅니다. 비밀번호를 바꾸려면 --new-password 로, 옛 비밀번호를 SITE_OLD_PASSWORD 에 넣어 실행하세요.")
        else:
            op = os.environ.get("SITE_OLD_PASSWORD")
            if not op: sys.exit("--new-password 는 SITE_OLD_PASSWORD(옛 비밀번호)가 필요합니다.")
            old_key = derive(op, unb64(old["kdf"]["salt"]))
        for s in old.get("seasons", []):
            f = out / "seasons" / f"{s['label']}.json"
            if f.exists():
                keep[s["label"]] = dec(old_key, json.loads(f.read_text(encoding="utf-8")))

    salt = unb64(old["kdf"]["salt"]) if (old and not args.new_password) else os.urandom(16)
    key = derive(password, salt)

    stages, seasons = read_sheet(args.xlsx)
    products = read_products(args.products) if args.products else []
    for it in products:
        label = it["season"] or "시즌 미지정"
        # 진행현황 시트에 같은 품번이 있으면 제품 파일이 우선
        seasons[label] = [x for x in seasons.get(label, []) if x.get("code") != it["code"]] + [it]
    if products: print("제품 파일:", ", ".join(f"{p['code']}({p['season']})" for p in products))
    now = dt.datetime.now(KST).strftime("%Y-%m-%d %H:%M")
    payloads = dict(keep)
    for label, items in seasons.items():
        payloads[label] = {"label": label, "updated": now, "stages": stages, "sheetUrl": args.sheet_url, "items": items}

    entries = []
    for label in sorted(payloads, key=season_key, reverse=True):
        p = payloads[label]
        (out / "seasons" / f"{label}.json").write_text(json.dumps(enc(key, p), ensure_ascii=False), encoding="utf-8")
        entries.append({"label": label, "file": f"data/seasons/{label}.json", "updated": p["updated"]})
        if args.plain:
            pd = Path("data_plain"); pd.mkdir(exist_ok=True)
            (pd / f"{label}.json").write_text(json.dumps(p, ensure_ascii=False, indent=1), encoding="utf-8")

    manifest = {"v": 1, "updated": now, "kdf": {"name": "PBKDF2-SHA256", "salt": b64(salt), "iter": ITER},
                "check": enc(key, CHECK_TEXT), "seasons": entries}
    man_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8")

    for label in sorted(seasons, key=season_key, reverse=True):
        print(f"{label}: {len(seasons[label])}건 갱신")
    kept = [l for l in keep if l not in seasons]
    if kept: print("유지(시트에 없음):", ", ".join(kept))
    print("완료:", now)


if __name__ == "__main__":
    main()
